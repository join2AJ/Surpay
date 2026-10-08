"""Import a surplus list someone downloaded by hand (CSV or XLSX).

Many counties publish spreadsheets, or send one in reply to a public records request.
Columns are mapped by name, e.g.:

    python -m surpay.cli import-file list.xlsx --state FL --county Lee \
        --sale-type tax_sale --name "Owner Name" --amount "Surplus" \
        --reference "Case Number" --address "Mailing Address" --sale-date "Sale Date"
"""

import csv
from datetime import date, datetime
from pathlib import Path

from .base import RecordIn, parse_money

_DATE_FORMATS = ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%m-%d-%Y")


def _to_date(value) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Unrecognized date: {value!r}")


def _to_cents(value) -> int:
    if isinstance(value, (int, float)):
        return round(value * 100)
    return parse_money(str(value))


def read_rows(path: Path) -> list[dict]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
        return [dict(zip(header, r)) for r in rows if any(c not in (None, "") for c in r)]
    with path.open(newline="", encoding="utf-8-sig") as f:
        return [{k.strip(): v for k, v in row.items()} for row in csv.DictReader(f)]


def records_from_file(
    path: Path,
    *,
    state: str,
    county: str,
    sale_type: str,
    name_col: str,
    amount_col: str,
    reference_col: str,
    address_col: str | None = None,
    sale_date_col: str | None = None,
    source_url: str = "",
) -> list[RecordIn]:
    out = []
    for i, row in enumerate(read_rows(path), start=2):
        name = str(row.get(name_col) or "").strip()
        amount = row.get(amount_col)
        ref = str(row.get(reference_col) or "").strip()
        if not name or amount in (None, "") or not ref:
            continue
        try:
            cents = _to_cents(amount)
            sale = _to_date(row.get(sale_date_col)) if sale_date_col else None
        except ValueError as e:
            raise ValueError(f"{path.name} row {i}: {e}") from e
        if cents <= 0:
            continue
        out.append(
            RecordIn(
                source_key=f"{ref}|{sale.isoformat() if sale else ''}",
                state=state.upper(),
                county=county,
                sale_type=sale_type,
                reference=ref,
                owner_name=name,
                owner_address=str(row.get(address_col) or "").strip() if address_col else "",
                amount_cents=cents,
                sale_date=sale,
                source_url=source_url,
                raw={k: str(v) for k, v in row.items()},
            )
        )
    return out
