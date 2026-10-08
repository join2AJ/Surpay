import io
from datetime import date
from pathlib import Path

import docx

from surpay.scrapers.adams_oh import document_lines, find_list_url, parse_lines
from surpay.scrapers.base import parse_money
from surpay.scrapers.file_import import records_from_file

FIXTURES = Path(__file__).parent / "fixtures"

# Synthetic entries in the exact layout of the Adams County list (names are made up).
ADAMS_LINES = [
    "LIST OF SURPLUS FUNDS HELD IN ADAMS COUNTY, OHIO, DUE TO OVERAGES FROM TAX SALES:",
    "$4,686.31\tJane Testperson, 850 Example Lane, Peebles, OH 45660",
    "Date of Sale: \t6/17/24\tPayment received: 6/24/24\tParcel # 068-00-00-045.019",
    "$19,732.60\tRobert Sample, 6063 State Route 73, Peebles, OH 45660",
    "Date of Sale:\t7/1/24\t\tPayment received:  7/29/24\tParcel # 029-00-00-049.000",
    "\t\t\t\t\t\t\t\tParcel # 029.00-00-049.001",
    "$10,812.25        Alan Placeholder, 2937 Blacks Run, Lynx, OH 45650",
    "Date of Sale:\t11/4/24\tPayment received:  11/4/24\tParcel #148-00-00-042.003",
    "Updated 12/8/25",
]


def test_parse_money():
    assert parse_money("$48,009.13") == 4800913
    assert parse_money("$5") == 500
    assert parse_money("1,000.5") == 100050


def test_find_list_url_on_real_page_snapshot():
    html = (FIXTURES / "adams_oh_surplus_page.html").read_text()
    assert find_list_url(html) == "https://adamsoh-auditor.schneidergis.com/media/14046/surplus-funds-list.docx"


def test_parse_adams_lines():
    recs = parse_lines(ADAMS_LINES, "https://example/list.docx")
    assert [r.owner_name for r in recs] == ["Jane Testperson", "Robert Sample", "Alan Placeholder"]
    assert recs[0].amount_cents == 468631
    assert recs[0].owner_address == "850 Example Lane, Peebles, OH 45660"
    assert recs[0].sale_date == date(2024, 6, 17)
    assert recs[1].reference == "Parcel 029-00-00-049.000, 029.00-00-049.001"
    assert recs[2].sale_date == date(2024, 11, 4)
    assert len({r.source_key for r in recs}) == 3
    assert all(r.state == "OH" and r.county == "Adams" for r in recs)


def test_parse_adams_docx_roundtrip():
    d = docx.Document()
    for line in ADAMS_LINES:
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    recs = parse_lines(document_lines(buf.getvalue(), "x.docx"), "x.docx")
    assert len(recs) == 3
    assert sum(r.amount_cents for r in recs) == 468631 + 1973260 + 1081225


def test_entry_without_parcel_is_skipped():
    recs = parse_lines(["$100.00\tNo Parcel, 1 A St", "Date of Sale: 1/1/24"], "u")
    assert recs == []


def test_file_import_csv(tmp_path):
    p = tmp_path / "list.csv"
    p.write_text(
        "Case Number,Owner Name,Mailing Address,Surplus,Sale Date\n"
        "2023-CA-001,\"DOE, JOHN\",\"12 Elm St, Fort Myers, FL\",\"$12,500.00\",03/15/2023\n"
        "2023-CA-002,EMPTY AMOUNT,,,\n"
        "2023-CA-003,ZERO,,0,\n"
    )
    recs = records_from_file(p, state="fl", county="Lee", sale_type="tax_sale",
                             name_col="Owner Name", amount_col="Surplus", reference_col="Case Number",
                             address_col="Mailing Address", sale_date_col="Sale Date")
    assert len(recs) == 1
    r = recs[0]
    assert (r.state, r.amount_cents, r.sale_date) == ("FL", 1250000, date(2023, 3, 15))


def test_file_import_xlsx(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Parcel", "Name", "Excess"])
    ws.append(["12-34", "Mary Example", 2500.5])
    p = tmp_path / "list.xlsx"
    wb.save(p)
    recs = records_from_file(p, state="OH", county="Test", sale_type="tax_sale",
                             name_col="Name", amount_col="Excess", reference_col="Parcel")
    assert recs[0].amount_cents == 250050
